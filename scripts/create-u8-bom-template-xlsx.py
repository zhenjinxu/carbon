from __future__ import annotations

import argparse
import json
from copy import copy
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DETAIL_HEADERS = [
    "层级",
    "序号",
    "图号/ERP编码",
    "名称",
    "部件内数量",
    "整机数量",
    "材料/类别",
    "单重(kg)",
    "总重(kg)",
    "图纸类别",
    "规格/标准",
    "备注",
    "合并图页",
]

SUMMARY_HEADERS = ["ERP编码", "名称", "整机数量", "规格/标准", "材料/等级", "类别"]

TEMPLATE_DETAIL_SHEET = "整机BOM(多级)"
TEMPLATE_SUMMARY_SHEET = "标准件外购件汇总"


def decimal_or_none(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None


def excel_number(value):
    number = decimal_or_none(value)
    if number is None:
        return None
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def excel_text(value):
    if value in (None, ""):
        return None
    return str(value)


def category_for(item, has_child_bom=False):
    code = str(item.get("code") or "")
    name = str(item.get("name") or "")
    class_name = str(item.get("className") or "")
    is_purchased = bool(item.get("isPurchased"))
    is_manufactured = bool(item.get("isManufactured")) or has_child_bom

    if code.startswith("3201"):
        return "标准件"
    if code.startswith("3202"):
        return "外购件"
    if "螺栓" in name or "螺母" in name or "平垫" in name or "弹垫" in name or "螺钉" in name:
        return "标准件"
    if is_purchased and ("原料" in class_name or code.startswith("37")):
        return "原料"
    if is_purchased:
        return "外购件"
    if is_manufactured:
        return "自制件"
    return class_name or "未分类"


def material_or_class(item):
    return item.get("className") or item.get("classCode") or None


def spec_or_standard(item):
    return item.get("specification") or item.get("engineerFigureNo") or None


def total_weight(unit_weight, cumulative_qty):
    weight = decimal_or_none(unit_weight)
    qty = decimal_or_none(cumulative_qty)
    if weight is None or qty is None:
        return None
    return excel_number(weight * qty)


def thin_border():
    side = Side(style="thin", color="808080")
    return Border(left=side, right=side, top=side, bottom=side)


def safe_sheet_name(name, used):
    invalid = set('[]:*?/\\')
    cleaned = "".join("_" if ch in invalid else ch for ch in name).strip() or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    index = 2
    while candidate in used:
        suffix = f"_{index}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        index += 1
    used.add(candidate)
    return candidate


def setup_page(ws, widths):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for row in range(1, max(ws.max_row, 4) + 1):
        ws.row_dimensions[row].height = 20
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 26


def style_header_row(ws, max_col):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    border = thin_border()
    for cell in ws[3][:max_col]:
        cell.fill = copy(fill)
        cell.font = copy(font)
        cell.border = copy(border)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_title_rows(ws, max_col):
    border = thin_border()
    for row in (1, 2):
        for cell in ws[row][:max_col]:
            cell.border = copy(border)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws[1][0].font = Font(bold=True, size=14)
    ws[2][0].font = Font(italic=True, color="666666")


def style_detail_rows(ws, first_row, last_row):
    border = thin_border()
    top_level_fill = PatternFill("solid", fgColor="EDF2FA")
    max_outline_level = 0
    for row in range(first_row, last_row + 1):
        level = ws.cell(row=row, column=1).value
        try:
            level_number = int(level)
        except (TypeError, ValueError):
            level_number = 0
        if level_number > 0:
            outline_level = min(level_number, 7)
            ws.row_dimensions[row].outlineLevel = outline_level
            max_outline_level = max(max_outline_level, outline_level)
        is_top_level = level_number == 1
        for col in range(1, 14):
            cell = ws.cell(row=row, column=col)
            cell.border = copy(border)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if is_top_level:
                cell.fill = copy(top_level_fill)
                cell.font = Font(bold=True)
        ws.cell(row=row, column=4).alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
            indent=max(0, level_number - 1),
        )
        for col in (1, 2, 3):
            ws.cell(row=row, column=col).number_format = "@"
        for col in (5, 6, 8, 9):
            ws.cell(row=row, column=col).number_format = "0.######"
    ws.sheet_format.outlineLevelRow = max_outline_level


def add_detail_sheet(wb, root_payload, sheet_name):
    root = root_payload["root"]
    rows = root_payload["rows"]
    ws = wb.create_sheet(sheet_name)
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=13)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=13)
    ws.cell(row=1, column=1).value = f'{root.get("code")} {root.get("name")} 整机BOM表'
    ws.cell(row=2, column=1).value = (
        f'U8已审核BOM导出 ｜ BomId {root_payload.get("selectedBom", {}).get("bomId")} '
        f'｜ Version {root_payload.get("selectedBom", {}).get("version")} '
        f'｜ As of {root_payload.get("asOfDate")}'
    )
    for col, header in enumerate(DETAIL_HEADERS, 1):
        ws.cell(row=3, column=col).value = header

    root_weight = excel_number(root.get("unitWeight"))
    root_values = [
        "整机",
        None,
        excel_text(root.get("code")),
        root.get("name"),
        None,
        1,
        material_or_class(root),
        root_weight,
        root_weight,
        category_for(root, True),
        spec_or_standard(root),
        f'BomId {root_payload.get("selectedBom", {}).get("bomId")}',
        None,
    ]
    for col, value in enumerate(root_values, 1):
        ws.cell(row=4, column=col).value = value

    for index, row in enumerate(rows, 5):
        item = row["item"]
        cumulative = row.get("cumulativeQuantity")
        values = [
            excel_text(row.get("level") or 0),
            excel_text(row.get("siblingSeq") or 0),
            excel_text(item.get("code")),
            item.get("name"),
            excel_number(row.get("directQuantity")),
            excel_number(cumulative),
            material_or_class(item),
            excel_number(item.get("unitWeight")),
            total_weight(item.get("unitWeight"), cumulative),
            category_for(item, row.get("hasChildBom")),
            spec_or_standard(item),
            row.get("edge", {}).get("remark") if row.get("edge") else None,
            None,
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=index, column=col).value = value

    widths = {
        "A": 6,
        "B": 6,
        "C": 16,
        "D": 30,
        "E": 10,
        "F": 9,
        "G": 18,
        "H": 9,
        "I": 9,
        "J": 9,
        "K": 26,
        "L": 30,
        "M": 9,
    }
    setup_page(ws, widths)
    style_title_rows(ws, 13)
    style_header_row(ws, 13)

    root_fill = PatternFill("solid", fgColor="D9E2F3")
    for cell in ws[4][:13]:
        cell.fill = copy(root_fill)
        cell.font = Font(bold=True)
        cell.border = thin_border()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row=4, column=4).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.cell(row=4, column=3).number_format = "@"
    ws.cell(row=4, column=6).number_format = "0.######"
    if rows:
        style_detail_rows(ws, 5, 4 + len(rows))
    return ws


def summary_rows(root_payload):
    grouped = {}
    for row in root_payload["rows"]:
        item = row["item"]
        category = category_for(item, row.get("hasChildBom"))
        if category not in {"标准件", "外购件", "原料"} and not item.get("isPurchased"):
            continue
        code = excel_text(item.get("code"))
        key = (code, item.get("name"), spec_or_standard(item), material_or_class(item), category)
        grouped.setdefault(key, Decimal("0"))
        grouped[key] += decimal_or_none(row.get("cumulativeQuantity")) or Decimal("0")
    result = []
    for (code, name, spec, material, category), qty in grouped.items():
        result.append([code, name, excel_number(qty), spec, material, category])
    result.sort(key=lambda values: str(values[0]))
    return result


def add_summary_sheet(wb, root_payload, sheet_name):
    root = root_payload["root"]
    ws = wb.create_sheet(sheet_name)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    ws.cell(row=1, column=1).value = f'{root.get("code")} 标准件 / 外购件整机汇总（按ERP编码聚合，含各级装配用量）'
    for col, header in enumerate(SUMMARY_HEADERS, 1):
        ws.cell(row=3, column=col).value = header
    rows = summary_rows(root_payload)
    for row_index, values in enumerate(rows, 4):
        for col, value in enumerate(values, 1):
            ws.cell(row=row_index, column=col).value = value
    setup_page(ws, {"A": 18, "B": 30, "C": 10, "D": 26, "E": 18, "F": 12})
    style_title_rows(ws, 6)
    style_header_row(ws, 6)
    border = thin_border()
    for row in range(4, 4 + len(rows)):
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            cell.border = copy(border)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row=row, column=4).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row=row, column=1).number_format = "@"
        ws.cell(row=row, column=3).number_format = "0.######"
    return ws


def workbook_for_roots(payload, roots):
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    single = len(roots) == 1
    for root_payload in roots:
        root_payload = {**root_payload, "asOfDate": payload.get("asOfDate")}
        code = root_payload["root"].get("code")
        detail_name = TEMPLATE_DETAIL_SHEET if single else safe_sheet_name(f"{code}_多级BOM", used)
        if single:
            used.add(detail_name)
        add_detail_sheet(wb, root_payload, detail_name)
        summary_name = TEMPLATE_SUMMARY_SHEET if single else safe_sheet_name(f"{code}_汇总", used)
        if single:
            used.add(summary_name)
        add_summary_sheet(wb, root_payload, summary_name)
    return wb


def validate_workbook(path):
    wb = load_workbook(path, data_only=False)
    for ws in wb.worksheets:
        if ws.max_row < 4 or ws.max_column < 6:
            raise ValueError(f"{ws.title} did not render expected rows/columns")
        if ws.cell(row=1, column=1).value in (None, ""):
            raise ValueError(f"{ws.title} missing title")
        if ws.cell(row=3, column=1).value in (None, ""):
            raise ValueError(f"{ws.title} missing header")
    return {"sheets": wb.sheetnames, "dimensions": {ws.title: ws.calculate_dimension() for ws in wb.worksheets}}


def main():
    parser = argparse.ArgumentParser(description="Create U8 BOM Excel files from export-u8-bom-template JSON output.")
    parser.add_argument("--input", required=True, help="JSON file generated by export-u8-bom-template.cjs")
    parser.add_argument("--output", help="Single .xlsx output path. Multiple roots become one workbook with sheet pairs.")
    parser.add_argument("--output-dir", help="Directory for one workbook per root. Default when JSON has multiple roots.")
    args = parser.parse_args()

    input_path = Path(args.input).resolve()
    payload = json.loads(input_path.read_text(encoding="utf-8"))
    roots = payload.get("roots") or []
    if not roots:
        raise SystemExit("No roots found in input JSON.")

    outputs = []
    if args.output:
        output_path = Path(args.output).resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb = workbook_for_roots(payload, roots)
        wb.save(output_path)
        outputs.append({"path": str(output_path), **validate_workbook(output_path)})
    else:
        output_dir = Path(args.output_dir or input_path.parent / "u8-bom-template-xlsx").resolve()
        output_dir.mkdir(parents=True, exist_ok=True)
        for root_payload in roots:
            code = root_payload["root"].get("code")
            output_path = output_dir / f"u8-bom-{code}-template.xlsx"
            wb = workbook_for_roots(payload, [root_payload])
            wb.save(output_path)
            outputs.append({"path": str(output_path), **validate_workbook(output_path)})

    print(json.dumps({"input": str(input_path), "outputs": outputs}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
